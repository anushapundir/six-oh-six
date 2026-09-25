"""Render an Analysis as a technical accounting memo (.docx)."""

import io
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

from sixohsix.schema import KINDS, Analysis, Case, Timing
from sixohsix.score import cited_ids

FONT = "Calibri"
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x1D, 0x4E, 0x89)
TIMING = {Timing.OVER_TIME: "Over time", Timing.POINT_IN_TIME: "Point in time"}


def style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = FONT, Pt(10.5), INK
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 20), ("Heading 1", 13), ("Heading 2", 11)):
        s = doc.styles[name]
        s.font.name, s.font.size, s.font.color.rgb, s.font.bold = FONT, Pt(size), ACCENT if name != "Title" else INK, True
        s.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 8)


def table(doc: Document, header: list[str] | None, rows: list[list[str]]):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1" if header else "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(([header] if header else []) + rows):
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.bold = bool(header) and i == 0
    return t


def render(case: Case, analysis: Analysis) -> bytes:
    doc = Document()
    style(doc)
    doc.add_paragraph("ASC 606 Revenue Recognition Memo", style="Title")

    note = doc.add_paragraph()
    run = note.add_run(
        "Draft for review. Prepared by an automated analysis and not reviewed by a qualified accountant. "
        "Verify every conclusion against the contract and the codification before relying on it."
    )
    run.italic, run.font.color.rgb, run.font.size = True, MUTED, Pt(9.5)

    header = table(doc, None, [
        ["Entity", analysis.reporting_entity],
        ["Customer", case.customer],
        ["Contract", f"{case.title} ({case.contract_type})"],
        ["Date", f"{date.today():%B} {date.today().day}, {date.today().year}"],
        ["Prepared by", "six-oh-six agent (draft)"],
    ])
    for row in header.rows:
        row.cells[0].paragraphs[0].runs[0].bold = True

    doc.add_heading("Purpose", 1)
    doc.add_paragraph(
        f"To document how {analysis.reporting_entity} should recognize revenue under ASC 606 for its contract with "
        f"{case.customer}: the performance obligations, the transaction price, and when each obligation is satisfied."
    )

    doc.add_heading("Background", 1)
    doc.add_paragraph(analysis.summary)
    doc.add_paragraph(f"Source: {case.source}. The contract has {len(case.clauses)} clauses; clause ids are cited in brackets.")

    doc.add_heading("Analysis", 1)
    doc.add_heading("Step 1. Identify the contract", 2)
    doc.add_paragraph(
        "The agreement is treated as a contract within the scope of ASC 606 (606-10-25-1). Any doubt about "
        "enforceability or collectability is listed under Open questions."
    )
    doc.add_heading("Step 2. Identify the performance obligations", 2)
    table(doc, ["Obligation", "Timing", "Basis", "Clauses", "Guidance"], [
        [f"{KINDS[o.kind].label}\n{o.description}", TIMING[o.timing], o.rationale, ", ".join(o.clauses), ", ".join(o.guidance)]
        for o in analysis.obligations
    ] or [["None identified", "", "", "", ""]])
    doc.add_heading("Step 3. Determine the transaction price", 2)
    doc.add_paragraph(analysis.consideration.summary)
    doc.add_heading("Step 4. Allocate the transaction price", 2)
    doc.add_paragraph(
        "Fixed consideration is allocated to the obligations above on a relative standalone selling price basis "
        "(606-10-32-28 to 32-31). Sales- or usage-based royalties that qualify for the exception are allocated to the "
        "license they relate to."
    )
    doc.add_heading("Step 5. Recognize revenue as obligations are satisfied", 2)
    for o in analysis.obligations:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{KINDS[o.kind].label}: ").bold = True
        p.add_run(f"{TIMING[o.timing].lower()}. {o.rationale}")

    c = analysis.consideration
    doc.add_heading("Consideration", 1)
    table(doc, None, [
        ["Variable consideration", "Yes" if c.variable else "No"],
        ["Sales- or usage-based royalty exception (606-10-55-65)", "Applies" if c.royalty_exception else "Does not apply"],
        ["Clauses", ", ".join(c.clauses)],
    ])
    doc.add_paragraph()
    doc.add_paragraph(c.rationale)

    doc.add_heading("Conclusion", 1)
    over = [KINDS[o.kind].label.lower() for o in analysis.obligations if o.timing == Timing.OVER_TIME]
    point = [KINDS[o.kind].label.lower() for o in analysis.obligations if o.timing == Timing.POINT_IN_TIME]
    parts = []
    if point:
        parts.append(f"recognize revenue for {'; '.join(point)} at the point control transfers")
    if over:
        parts.append(f"recognize revenue for {'; '.join(over)} over time as it performs")
    conclusion = f"{analysis.reporting_entity} should {' and '.join(parts) or 'reassess the contract'}."
    if c.royalty_exception:
        conclusion += " The royalty is recognized as the customer's underlying sales or usage occur, not estimated up front."
    elif c.variable:
        conclusion += " Variable consideration is estimated and constrained under 606-10-32-11."
    doc.add_paragraph(conclusion)

    doc.add_heading("Open questions", 1)
    for q in analysis.open_questions or ["None raised."]:
        doc.add_paragraph(q, style="List Number")

    doc.add_heading("Appendix: cited clauses", 1)
    by_id = {cl.id: cl for cl in case.clauses}
    cited = set(cited_ids(analysis))
    for cl in case.clauses:
        if cl.id in cited:
            doc.add_paragraph().add_run(f"[{cl.id}] {cl.heading}").bold = True
            doc.add_paragraph(cl.text)
    missing = sorted(cited - by_id.keys())
    if missing:
        doc.add_paragraph().add_run(f"Cited but not found in the contract: {', '.join(missing)}").italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
